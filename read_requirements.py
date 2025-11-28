import openpyxl
import os
import glob
import json

# Find XLSX files
xlsx_files = glob.glob("*.xlsx")
if xlsx_files:
    file_path = xlsx_files[0]
    print(f"Found: {file_path}")
    print(f"Full path: {os.path.abspath(file_path)}")
    
    try:
        wb = openpyxl.load_workbook(file_path)
        print(f"\nSheets: {wb.sheetnames}\n")
        
        # Read all sheets
        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            print(f"\n{'='*80}")
            print(f"Sheet: {sheet_name}")
            print(f"{'='*80}")
            
            # Get max row and column with data
            max_row = ws.max_row
            max_col = ws.max_column
            
            print(f"Rows: {max_row}, Columns: {max_col}\n")
            
            # Print all data
            for row_idx, row in enumerate(ws.iter_rows(min_row=1, max_row=max_row, values_only=True), 1):
                if any(row):  # Only print non-empty rows
                    print(f"Row {row_idx}: {row}")
        
        wb.close()
        
    except Exception as e:
        print(f"Error: {e}")
        import traceback
        traceback.print_exc()
else:
    print("No XLSX files found")
